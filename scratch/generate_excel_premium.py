"""
Excel Premium de Pilotage - Elke Berr.
Inclus Menus deroulants, Dashboard de stats et structuration par composant.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUTPUT = "c:/Users/lyies/Desktop/siteelk/wireframes_pack/Elke_Berr_Pilotage_Premium.xlsx"

def create_premium_excel():
    wb = Workbook()
    
    # --- FEUILLE 1 : RÉSUMÉ & DASHBOARD ---
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash["B2"] = "PROJET : ELKE BERR SITE WEB V2"
    ws_dash["B2"].font = Font(size=14, bold=True)
    ws_dash["B4"] = "Statistiques de Validation"
    ws_dash["B4"].font = Font(bold=True)
    
    # Table de stats (formules pour plus tard si besoin)
    stats = [
        ["Total Sections", 6],
        ["Validées", 0],
        ["À modifier", 0],
        ["Refusées", 0],
        ["Taux de validation", "0%"]
    ]
    for r_idx, (label, val) in enumerate(stats, 6):
        ws_dash.cell(row=r_idx, column=2, value=label)
        ws_dash.cell(row=r_idx, column=3, value=val)

    # --- FEUILLE 2 : AUDIT DÉTAILLÉ ---
    ws_audit = wb.create_sheet("Audit de Détail")
    headers = [
        "ID", "Section", "Composant", "Description", 
        "Statut", "Priorité", "Commentaire Client", "Action Plan"
    ]
    ws_audit.append(headers)
    
    # Data
    data = [
        ["01", "Hero", "Navigation", "Menu principal (Maison, Collections...)", "A VALIDER", "Haute", "", ""],
        ["01", "Hero", "Accroche", "Titre 'La ou la pierre dicte l'oeuvre'", "A VALIDER", "Moyenne", "", ""],
        ["02", "Heritage", "Portrait", "Photo Paul Pogba", "A VALIDER", "Basse", "", ""],
        ["03", "Collections", "Grille", "Affichage des 3 bijoux signature", "A VALIDER", "Haute", "", ""],
        # ... etc
    ]
    for r in data:
        ws_audit.append(r)

    # Menus deroulants (Validation de donnees)
    dv_status = DataValidation(type="list", formula1='"A VALIDER,VALIDÉ ✅,À MODIFIER 🛠️,REFUSÉ ❌"', allow_blank=True)
    dv_priority = DataValidation(type="list", formula1='"Haute,Moyenne,Basse"', allow_blank=True)
    
    ws_audit.add_data_validation(dv_status)
    ws_audit.add_data_validation(dv_priority)
    
    # Appliquer aux colonnes Statut (E) et Priorite (F)
    dv_status.add(f"E2:E100")
    dv_priority.add(f"F2:F100")

    # Styling Audit
    header_fill = PatternFill(start_color="1a1a1a", end_color="1a1a1a", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)
    for cell in ws_audit[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center")

    ws_audit.column_dimensions['B'].width = 20
    ws_audit.column_dimensions['C'].width = 25
    ws_audit.column_dimensions['D'].width = 40
    ws_audit.column_dimensions['G'].width = 50
    
    wb.save(OUTPUT)
    print(f"Excel Premium genere : {OUTPUT}")

if __name__ == "__main__":
    create_premium_excel()
